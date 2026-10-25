"""Local desktop data workbench. Run: python data_workbench.py"""
import io
import json
import os
import re
import sqlite3
import subprocess
import sys
import tempfile
import zlib
from datetime import datetime
from pathlib import Path

import pandas as pd
from defusedxml import ElementTree as SafeET
from docx import Document
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg
from matplotlib.figure import Figure
from PySide6.QtCore import QAbstractTableModel, QModelIndex, Qt
from PySide6.QtWidgets import (
    QApplication, QComboBox, QFileDialog, QFormLayout, QHBoxLayout, QInputDialog,
    QLabel, QLineEdit, QMainWindow, QMessageBox, QPushButton, QSplitter,
    QTableView, QTextEdit, QVBoxLayout, QWidget,
)

APP_DIR = Path(os.getenv("LOCALAPPDATA", str(Path.home()))) / "DataWorkbench"
APP_DIR.mkdir(parents=True, exist_ok=True)
WORKSPACE = APP_DIR / "workspace.sqlite3"
PREVIEW_ROWS = 500


def unique_columns(frame):
    frame = frame.copy()
    used = set()
    names = []
    for i, col in enumerate(frame.columns):
        base = str(col).strip() or f"column_{i + 1}"
        name, counter = base, 2
        while name.casefold() in used:
            name = f"{base}_{counter}"
            counter += 1
        used.add(name.casefold())
        names.append(name)
    frame.columns = names
    return frame.reset_index(drop=True)


def text_frame(text):
    return pd.DataFrame({"line": text.splitlines() or [""]})


def parse_xml(path):
    root = SafeET.parse(path).getroot()
    elements = list(root)
    if not elements:
        elements = [root]
    rows = []
    for item in elements:
        row = {"tag": item.tag}
        row.update({"@" + k: v for k, v in item.attrib.items()})
        for child in item:
            key = child.tag
            value = " ".join(child.itertext()).strip()
            row[key] = (row[key] + " | " + value) if key in row else value
        if not list(item):
            row["text"] = (item.text or "").strip()
        rows.append(row)
    return pd.DataFrame(rows)


def sqlite_tables(connection):
    rows = connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name").fetchall()
    return [row[0] for row in rows]


def read_file(path, parent):
    suffix = path.suffix.lower()
    if suffix == ".csv":
        try:
            return pd.read_csv(path, encoding="utf-8-sig", sep=None, engine="python")
        except UnicodeDecodeError:
            return pd.read_csv(path, encoding="cp1252", sep=None, engine="python")
    if suffix in (".xls", ".xlsx"):
        book = pd.ExcelFile(path)
        sheet, ok = QInputDialog.getItem(parent, "Choose sheet", "Worksheet:", book.sheet_names, 0, False)
        if not ok:
            return None
        return pd.read_excel(book, sheet_name=sheet)
    if suffix == ".txt":
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            text = path.read_text(encoding="cp1252")
        first = "\n".join(text.splitlines()[:10])
        if first and any(c in first for c in ("\t", ",", ";", "|")):
            try:
                return pd.read_csv(io.StringIO(text), sep=None, engine="python")
            except (ValueError, pd.errors.ParserError):
                pass
        return text_frame(text)
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(data, list):
            return pd.json_normalize(data) if data else pd.DataFrame()
        if isinstance(data, dict):
            for value in data.values():
                if isinstance(value, list) and value and isinstance(value[0], dict):
                    return pd.json_normalize(value)
            return pd.json_normalize(data)
        return pd.DataFrame({"value": [data]})
    if suffix == ".xml":
        return parse_xml(path)
    if suffix == ".pdf":
        import pdfplumber
        lines = []
        with pdfplumber.open(path) as pdf:
            for page_no, page in enumerate(pdf.pages, 1):
                for line in (page.extract_text() or "").splitlines():
                    lines.append({"page": page_no, "line": line})
        return pd.DataFrame(lines, columns=["page", "line"])
    if suffix in (".doc", ".docx"):
        if suffix == ".doc":
            with tempfile.TemporaryDirectory() as temp:
                result = subprocess.run(
                    ["soffice", "-env:UserInstallation=file:///" + temp.replace("\\", "/") + "/lo_profile", "--headless", "--convert-to", "docx", "--outdir", temp, str(path)],
                    capture_output=True, text=True, timeout=120,
                )
                converted = Path(temp) / (path.stem + ".docx")
                if result.returncode != 0 or not converted.exists():
                    raise ValueError(".doc requires LibreOffice (soffice) installed and available on PATH")
                return read_docx(converted)
        return read_docx(path)
    if suffix in (".db", ".sqlite", ".sqlite3", ".sql"):
        if suffix == ".sql":
            answer = QMessageBox.question(parent, "Run SQL dump?", "Only import SQL dumps you trust. The script runs in a temporary in-memory SQLite database. Continue?")
            if answer != QMessageBox.StandardButton.Yes:
                return None
            connection = sqlite3.connect(":memory:")
            # Prevent dump scripts from attaching databases or writing via VACUUM INTO.
            forbidden = {sqlite3.SQLITE_ATTACH, sqlite3.SQLITE_DETACH, sqlite3.SQLITE_PRAGMA}
            if hasattr(sqlite3, "SQLITE_VACUUM"):
                forbidden.add(sqlite3.SQLITE_VACUUM)
            connection.set_authorizer(lambda action, *_: sqlite3.SQLITE_DENY if action in forbidden else sqlite3.SQLITE_OK)
            try:
                connection.executescript(path.read_text(encoding="utf-8-sig"))
                return choose_table(connection, parent)
            finally:
                connection.close()
        connection = sqlite3.connect(path.resolve().as_uri() + "?mode=ro", uri=True)
        try:
            return choose_table(connection, parent)
        finally:
            connection.close()
    raise ValueError(f"Unsupported extension: {suffix}")


def read_docx(path):
    document = Document(path)
    rows = [{"type": "paragraph", "text": p.text} for p in document.paragraphs if p.text.strip()]
    for number, table in enumerate(document.tables, 1):
        for row in table.rows:
            rows.append({"type": f"table {number}", "text": " | ".join(cell.text for cell in row.cells)})
    return pd.DataFrame(rows, columns=["type", "text"])


def choose_table(connection, parent):
    tables = sqlite_tables(connection)
    if not tables:
        raise ValueError("No data tables found in the SQLite database")
    table, ok = QInputDialog.getItem(parent, "Choose table", "Table:", tables, 0, False)
    if not ok:
        return None
    return pd.read_sql_query('SELECT * FROM "' + table.replace('"', '""') + '"', connection)


def serialize(frame):
    return zlib.compress(frame.to_json(orient="split", date_format="iso", default_handler=str).encode("utf-8"))


def deserialize(blob):
    return pd.read_json(io.StringIO(zlib.decompress(blob).decode("utf-8")), orient="split")


def apply_conditions(frame, expression):
    if not expression.strip():
        raise ValueError("Enter a condition, such as [age] >= 18 AND [region] = West")
    mask = pd.Series(True, index=frame.index)
    for condition in re.split(r"\s+AND\s+", expression.strip(), flags=re.IGNORECASE):
        match = re.fullmatch(r"\s*\[([^]]+)\]\s*(>=|<=|!=|=|>|<|~)\s*(.*?)\s*", condition)
        if not match:
            raise ValueError("Use [column] operator value, joined by AND. Operators: = != > >= < <= ~")
        column, operator, value = match.groups()
        if column not in frame.columns:
            raise ValueError(f"Unknown column: {column}")
        series = frame[column]
        if operator in ("=", "!=") and value.lower() == "null":
            result = series.isna() if operator == "=" else series.notna()
        elif operator == "~":
            result = series.astype("string").str.contains(value, case=False, regex=False, na=False)
        elif operator in ("=", "!="):
            result = series.astype("string").str.casefold() == value.casefold()
            result = result.fillna(False)
            if operator == "!=":
                result = (~result) & series.notna()
        else:
            number = pd.to_numeric(value, errors="coerce")
            if pd.isna(number):
                raise ValueError("Comparison values for >, >=, <, <= must be numbers")
            numeric = pd.to_numeric(series, errors="coerce")
            result = {">": numeric.gt, ">=": numeric.ge, "<": numeric.lt, "<=": numeric.le}[operator](number).fillna(False)
        mask &= result.fillna(False)
    return frame.loc[mask].reset_index(drop=True)


class FrameModel(QAbstractTableModel):
    def __init__(self):
        super().__init__()
        self.frame = pd.DataFrame()

    def set_frame(self, frame):
        self.beginResetModel()
        self.frame = frame.head(PREVIEW_ROWS)
        self.endResetModel()

    def rowCount(self, parent=QModelIndex()):
        return len(self.frame)

    def columnCount(self, parent=QModelIndex()):
        return len(self.frame.columns)

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if role == Qt.ItemDataRole.DisplayRole and index.isValid():
            value = self.frame.iat[index.row(), index.column()]
            return "" if pd.isna(value) else str(value)
        return None

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if role == Qt.ItemDataRole.DisplayRole:
            return str(self.frame.columns[section]) if orientation == Qt.Orientation.Horizontal else str(section + 1)
        return None


class Workbench(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Data Workbench")
        self.resize(1250, 800)
        self.connection = sqlite3.connect(WORKSPACE)
        self.connection.execute("CREATE TABLE IF NOT EXISTS datasets (id INTEGER PRIMARY KEY, name TEXT NOT NULL, source TEXT NOT NULL, imported_at TEXT NOT NULL, original BLOB NOT NULL, current BLOB NOT NULL)")
        self.connection.commit()
        self.dataset_id = None
        self.original = pd.DataFrame()
        self.current = pd.DataFrame()
        self.model = FrameModel()
        self.build_ui()
        self.reload_datasets()

    def build_ui(self):
        main = QWidget()
        layout = QVBoxLayout(main)
        top = QHBoxLayout()
        self.datasets = QComboBox()
        self.datasets.currentIndexChanged.connect(self.load_dataset)
        top.addWidget(QLabel("Workspace dataset:"))
        top.addWidget(self.datasets, 1)
        for label, handler in (("Import file", self.import_file), ("Export", self.export_file), ("Reset to import", self.reset_data)):
            button = QPushButton(label)
            button.clicked.connect(handler)
            top.addWidget(button)
        layout.addLayout(top)
        filters = QHBoxLayout()
        self.prompt = QLineEdit()
        self.prompt.setPlaceholderText("Example: [age] >= 18 AND [city] ~ York")
        filters.addWidget(self.prompt, 1)
        button = QPushButton("Apply filter")
        button.clicked.connect(self.filter_data)
        filters.addWidget(button)
        self.presets = QComboBox()
        self.presets.addItems(["Choose preset", "Remove duplicates", "Drop rows with missing values", "Remove IQR outliers"])
        filters.addWidget(self.presets)
        button = QPushButton("Apply preset")
        button.clicked.connect(self.apply_preset)
        filters.addWidget(button)
        layout.addLayout(filters)
        chart = QHBoxLayout()
        self.chart_type = QComboBox()
        self.chart_type.addItems(["Histogram", "Bar (count by category)", "Line", "Scatter"])
        self.x_axis = QComboBox()
        self.y_axis = QComboBox()
        for label, widget in (("Chart", self.chart_type), ("X / value", self.x_axis), ("Y", self.y_axis)):
            chart.addWidget(QLabel(label))
            chart.addWidget(widget)
        button = QPushButton("Draw chart")
        button.clicked.connect(self.draw_chart)
        chart.addWidget(button)
        button = QPushButton("Save chart PNG")
        button.clicked.connect(self.save_chart)
        chart.addWidget(button)
        layout.addLayout(chart)
        splitter = QSplitter(Qt.Orientation.Vertical)
        table = QTableView()
        table.setModel(self.model)
        splitter.addWidget(table)
        self.figure = Figure(figsize=(9, 4))
        self.canvas = FigureCanvasQTAgg(self.figure)
        splitter.addWidget(self.canvas)
        splitter.setSizes([470, 250])
        layout.addWidget(splitter, 1)
        self.status = QLabel("Import a file to begin. Data stays on this PC.")
        layout.addWidget(self.status)
        self.log = QTextEdit()
        self.log.setReadOnly(True)
        self.log.setMaximumHeight(65)
        layout.addWidget(self.log)
        self.setCentralWidget(main)

    def error(self, exc):
        QMessageBox.critical(self, "Data Workbench", str(exc))

    def reload_datasets(self, select_id=None):
        rows = self.connection.execute("SELECT id, name, imported_at FROM datasets ORDER BY id DESC").fetchall()
        self.datasets.blockSignals(True)
        self.datasets.clear()
        for dataset_id, name, date in rows:
            self.datasets.addItem(f"{name} — {date}", dataset_id)
        if select_id is not None:
            index = self.datasets.findData(select_id)
            if index >= 0:
                self.datasets.setCurrentIndex(index)
        self.datasets.blockSignals(False)
        self.load_dataset()

    def load_dataset(self, *_):
        dataset_id = self.datasets.currentData()
        if dataset_id is None:
            return
        row = self.connection.execute("SELECT original, current FROM datasets WHERE id = ?", (dataset_id,)).fetchone()
        if row:
            self.dataset_id = dataset_id
            self.original, self.current = deserialize(row[0]), deserialize(row[1])
            self.refresh("Loaded saved dataset")

    def refresh(self, message):
        self.model.set_frame(self.current)
        self.x_axis.clear()
        self.y_axis.clear()
        columns = list(self.current.columns)
        self.x_axis.addItems(columns)
        self.y_axis.addItems(columns)
        self.status.setText(f"{message} | {len(self.current):,} rows × {len(columns)} columns | showing first {PREVIEW_ROWS} rows")
        self.log.append(f"{datetime.now():%H:%M:%S}  {message}")

    def persist(self, message):
        self.connection.execute("UPDATE datasets SET current = ? WHERE id = ?", (serialize(self.current), self.dataset_id))
        self.connection.commit()
        self.refresh(message)

    def import_file(self):
        filename, _ = QFileDialog.getOpenFileName(self, "Import data", "", "Data (*.csv *.xls *.xlsx *.txt *.pdf *.doc *.docx *.json *.xml *.sql *.db *.sqlite *.sqlite3)")
        if not filename:
            return
        try:
            frame = read_file(Path(filename), self)
            if frame is None:
                return
            frame = unique_columns(frame)
            if not len(frame.columns):
                raise ValueError("No columns could be extracted from this file")
            payload = serialize(frame)
            cursor = self.connection.execute("INSERT INTO datasets (name, source, imported_at, original, current) VALUES (?, ?, ?, ?, ?)", (Path(filename).name, filename, datetime.now().isoformat(timespec="seconds"), payload, payload))
            self.connection.commit()
            self.reload_datasets(cursor.lastrowid)
            self.refresh(f"Imported {Path(filename).name}")
        except Exception as exc:
            self.error(exc)

    def reset_data(self):
        if self.dataset_id is not None:
            self.current = self.original.copy()
            self.persist("Restored original imported data")

    def filter_data(self):
        if self.dataset_id is None:
            return
        try:
            result = apply_conditions(self.current, self.prompt.text())
            removed = len(self.current) - len(result)
            self.current = result
            self.persist(f"Applied filter; removed {removed:,} rows")
        except Exception as exc:
            self.error(exc)

    def apply_preset(self):
        if self.dataset_id is None:
            return
        try:
            choice = self.presets.currentIndex()
            if choice == 0:
                return
            old = len(self.current)
            if choice == 1:
                result = self.current.drop_duplicates()
            elif choice == 2:
                result = self.current.dropna()
            else:
                numeric = self.current.select_dtypes(include="number")
                if numeric.empty:
                    raise ValueError("IQR outlier removal needs numeric columns")
                q1, q3 = numeric.quantile(.25), numeric.quantile(.75)
                iqr = q3 - q1
                # Ignore constant columns; null values are not classified as outliers.
                outliers = ((numeric.lt(q1 - 1.5 * iqr)) | (numeric.gt(q3 + 1.5 * iqr))).fillna(False)
                result = self.current.loc[~outliers.any(axis=1)]
            self.current = result.reset_index(drop=True)
            self.persist(f"{self.presets.currentText()}; removed {old - len(self.current):,} rows")
        except Exception as exc:
            self.error(exc)

    def draw_chart(self):
        if self.dataset_id is None or self.current.empty:
            return
        try:
            x, y = self.x_axis.currentText(), self.y_axis.currentText()
            kind = self.chart_type.currentIndex()
            self.figure.clear()
            ax = self.figure.add_subplot(111)
            if kind == 0:
                values = pd.to_numeric(self.current[x], errors="coerce").dropna()
                if values.empty:
                    raise ValueError("Histogram requires numeric X values")
                ax.hist(values, bins=30)
                ax.set_xlabel(x)
            elif kind == 1:
                counts = self.current[x].astype("string").fillna("(missing)").value_counts().head(30)
                counts.plot.bar(ax=ax)
                ax.set_ylabel("Count (top 30 categories)")
                ax.set_xlabel(x)
            else:
                values = pd.to_numeric(self.current[y], errors="coerce")
                if not values.notna().any():
                    raise ValueError("Line and scatter charts require numeric Y values")
                if kind == 2:
                    ax.plot(self.current[x].astype(str), values)
                    if len(self.current) > 30:
                        ax.xaxis.set_major_locator(__import__("matplotlib").ticker.MaxNLocator(12))
                else:
                    xvalues = pd.to_numeric(self.current[x], errors="coerce")
                    if not xvalues.notna().any():
                        raise ValueError("Scatter requires numeric X values")
                    ax.scatter(xvalues, values, s=12)
                ax.set_xlabel(x)
                ax.set_ylabel(y)
            self.figure.tight_layout()
            self.canvas.draw()
            self.status.setText("Chart drawn from all current rows. Use Save chart PNG to export it.")
        except Exception as exc:
            self.error(exc)

    def save_chart(self):
        if not self.figure.axes:
            self.error(ValueError("Draw a chart first"))
            return
        filename, _ = QFileDialog.getSaveFileName(self, "Save chart", "chart.png", "PNG image (*.png)")
        if filename:
            self.figure.savefig(filename, dpi=150, bbox_inches="tight")

    def export_file(self):
        if self.dataset_id is None:
            return
        filename, _ = QFileDialog.getSaveFileName(self, "Export processed data", "processed.csv", "CSV (*.csv);;Excel (*.xlsx);;JSON (*.json);;Text (*.txt);;XML (*.xml);;PDF (*.pdf);;Word (*.docx);;SQLite (*.db);;SQL dump (*.sql)")
        if not filename:
            return
        try:
            frame = self.current
            path = Path(filename)
            suffix = path.suffix.lower()
            if suffix == ".csv":
                frame.to_csv(path, index=False, encoding="utf-8-sig")
            elif suffix == ".xlsx":
                frame.to_excel(path, index=False)
            elif suffix == ".json":
                frame.to_json(path, orient="records", force_ascii=False, indent=2, date_format="iso")
            elif suffix == ".txt":
                frame.to_csv(path, index=False, sep="\t", encoding="utf-8")
            elif suffix == ".xml":
                frame.to_xml(path, index=False, root_name="data", row_name="row")
            elif suffix == ".db":
                if path.exists():
                    path.unlink()
                with sqlite3.connect(path) as database:
                    frame.to_sql("data", database, index=False, if_exists="replace")
            elif suffix == ".sql":
                with tempfile.TemporaryDirectory() as temp:
                    with sqlite3.connect(Path(temp) / "export.db") as database:
                        frame.to_sql("data", database, index=False, if_exists="replace")
                        path.write_text("\n".join(database.iterdump()) + "\n", encoding="utf-8")
            elif suffix == ".docx":
                doc = Document()
                doc.add_heading("Processed data", 0)
                table = doc.add_table(rows=1, cols=len(frame.columns))
                for cell, col in zip(table.rows[0].cells, frame.columns):
                    cell.text = str(col)
                for row in frame.itertuples(index=False, name=None):
                    for cell, value in zip(table.add_row().cells, row):
                        cell.text = "" if pd.isna(value) else str(value)
                doc.save(path)
            elif suffix == ".pdf":
                from reportlab.lib import colors
                from reportlab.lib.pagesizes import landscape, letter
                from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
                from reportlab.lib.utils import simpleSplit
                if len(frame.columns) > 12:
                    raise ValueError("PDF export supports up to 12 columns; use CSV/XLSX for wide data")
                doc = SimpleDocTemplate(str(path), pagesize=landscape(letter))
                width = (landscape(letter)[0] - 72) / len(frame.columns)
                data = [list(map(str, frame.columns))]
                for row in frame.itertuples(index=False, name=None):
                    data.append(["\n".join(simpleSplit("" if pd.isna(v) else str(v), "Helvetica", 7, width - 8)[:3]) for v in row])
                table = Table(data, colWidths=[width] * len(frame.columns), repeatRows=1)
                table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .25, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), ("FONTSIZE", (0, 0), (-1, -1), 7)]))
                doc.build([table])
            else:
                raise ValueError("Choose a supported output extension")
            self.log.append(f"Exported {len(frame):,} rows to {path}")
        except Exception as exc:
            self.error(exc)

    def closeEvent(self, event):
        self.connection.close()
        super().closeEvent(event)


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = Workbench()
    window.show()
    sys.exit(app.exec())
